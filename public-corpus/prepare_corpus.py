#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
prepare_corpus.py (v2) — 把公开原始文本制备成 4 类格式源文件

v2 变更：
  - 去除小说源文件开头的目录块（连续 ≥8 行回目 → 判定为目录并删除）
  - 行首全角/半角空白规整（回目标题顶格，保证分章识别）
  - PDF 渲染时「回/章/卷」标题起新页（对齐实体书排版，OCR 后页首即回目）
"""
import json
import os
import re
import time

import fitz
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, "raw")
PREP = os.path.join(HERE, "prepared")
GOV = os.path.join(PREP, "gov")
os.makedirs(PREP, exist_ok=True)
os.makedirs(GOV, exist_ok=True)

A4 = fitz.paper_rect("a4")
MARGIN = 54
BODY = fitz.Rect(MARGIN, MARGIN, A4.width - MARGIN, A4.height - MARGIN)

HEADLINE = re.compile(
    r"^[ \t　]*(第\s*[\d０-９零一二三四五六七八九十百叁两]+\s*[章回卷]|卷\s*[一二三四五六七八九十]{1,3})\s")

SCAN_FONT = 13
SCAN_LEADING = 24
SCAN_CHARS_PER_PAGE = 860
TEXT_FONT = 10.5
TEXT_CHARS_PER_PAGE = 1750


def cjk_count(s):
    return sum(1 for c in s if "\u4e00" <= c <= "\u9fff")


def drop_toc(lines):
    """删除文件开头的目录块：前 300 行里最长的连续标题行段（≥8 行）。"""
    best_start, best_len, i = 0, 0, 0
    limit = min(len(lines), 300)
    while i < limit:
        if HEADLINE.match(lines[i]):
            j = i
            while j < limit and HEADLINE.match(lines[j]):
                j += 1
            if j - i > best_len:
                best_start, best_len = i, j - i
            i = j
        else:
            i += 1
    if best_len >= 8:
        return lines[:best_start] + lines[best_start + best_len:], best_len
    return lines, 0


def clean_source(path):
    s = open(path, encoding="utf-8").read().replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"\n{3,}", "\n\n", s).strip()
    lines = s.split("\n")
    lines, n_toc = drop_toc(lines)
    # 行首空白规整（保留段首一个全角空格供排版区分，标题行顶格）
    out = []
    for ln in lines:
        stripped = ln.lstrip(" \t　")
        if HEADLINE.match(stripped) or not stripped:
            out.append(stripped)
        else:
            out.append("　" + stripped)
    return "\n".join(out), n_toc


def split_pages_by_chars(text, chars_per_page, break_at_headings=True):
    lines, pages, cur, cur_len = text.split("\n"), [], [], 0
    for ln in lines:
        is_head = break_at_headings and HEADLINE.match(ln)
        add = len(ln) + 1
        if (cur_len + add > chars_per_page and cur) or (is_head and cur):
            pages.append("\n".join(cur))
            cur, cur_len = [], 0
        cur.append(ln)
        cur_len += add
    if cur:
        pages.append("\n".join(cur))
    return pages


def render_pdf(text, out_path, font_size, chars_per_page, leading):
    doc = fitz.open()
    pages = split_pages_by_chars(text, chars_per_page)
    for i, ptext in enumerate(pages, 1):
        pg = doc.new_page(width=A4.width, height=A4.height)
        rect = fitz.Rect(BODY.x0, BODY.y0, BODY.x1, BODY.y1 - 18)
        pg.insert_textbox(rect, ptext, fontname="china-s",
                          fontsize=font_size, lineheight=leading / font_size / 1.2,
                          align=fitz.TEXT_ALIGN_LEFT)
        pg.insert_text((BODY.x0, A4.height - 30), f"- {i} -",
                       fontname="china-s", fontsize=9)
    doc.save(out_path, deflate=True, garbage=3)
    doc.close()
    return len(pages)


def render_scanned_pdf(text, out_path, font_size, chars_per_page, leading, dpi=200):
    src = fitz.open()
    pages = split_pages_by_chars(text, chars_per_page)
    for i, ptext in enumerate(pages, 1):
        pg = src.new_page(width=A4.width, height=A4.height)
        rect = fitz.Rect(BODY.x0, BODY.y0, BODY.x1, BODY.y1 - 18)
        pg.insert_textbox(rect, ptext, fontname="china-s",
                          fontsize=font_size, lineheight=leading / font_size / 1.2)
        pg.insert_text((BODY.x0, A4.height - 30), f"- {i} -",
                       fontname="china-s", fontsize=9)
    out = fitz.open()
    t0 = time.time()
    for i, pg in enumerate(src):
        pix = pg.get_pixmap(dpi=dpi)
        np = out.new_page(width=A4.width, height=A4.height)
        np.insert_image(np.rect, pixmap=pix)
        if (i + 1) % 300 == 0:
            print(f"    rasterized {i+1}/{len(src)} pages  ({time.time()-t0:.0f}s)", flush=True)
    out.save(out_path, deflate=True)
    n = len(out)
    src.close()
    out.close()
    return n


def text_to_docx(text, out_path, title):
    doc = Document()
    doc.add_paragraph(title)
    for para in text.split("\n"):
        para = para.strip(" \t　")
        if not para:
            continue
        doc.add_paragraph(para)
    doc.save(out_path)


def main():
    report = {}

    hlm, n1 = clean_source(os.path.join(RAW, "hongloumeng.txt"))
    print(f"[prep] 红楼梦: 去目录 {n1} 行, cjk={cjk_count(hlm):,}")
    p1 = os.path.join(PREP, "红楼梦_文字版.pdf")
    if not os.path.exists(p1):
        n = render_pdf(hlm, p1, TEXT_FONT, TEXT_CHARS_PER_PAGE, 15)
        print(f"[pdf ] 红楼梦 -> {n} pages")
    report["红楼梦_文字版.pdf"] = {"kind": "pdf-text", "cjk": cjk_count(hlm)}

    for fname, out in [("liaozhai.txt", "聊斋志异_扫描版.pdf"),
                       ("fengshen_yanyi.txt", "封神演义_扫描版.pdf")]:
        t, nt = clean_source(os.path.join(RAW, fname))
        print(f"[prep] {out}: 去目录 {nt} 行, cjk={cjk_count(t):,}")
        p = os.path.join(PREP, out)
        if not os.path.exists(p):
            n = render_scanned_pdf(t, p, SCAN_FONT, SCAN_CHARS_PER_PAGE, SCAN_LEADING)
            print(f"[scan] {out} -> {n} pages (image-only)")
        report[out] = {"kind": "pdf-scan", "cjk": cjk_count(t)}

    for fname, out, title in [
        ("sanguo_yanyi.txt", "三国演义.docx", "三国演义"),
        ("shuihu_zhuan.txt", "水浒传.docx", "水浒传"),
        ("xiyouji.txt", "西游记.docx", "西游记"),
        ("jinghua_yuan.txt", "镜花缘.docx", "镜花缘"),
        ("dongzhou_lieguo.txt", "东周列国志.docx", "东周列国志"),
    ]:
        t, nt = clean_source(os.path.join(RAW, fname))
        p = os.path.join(PREP, out)
        if not os.path.exists(p):
            print(f"[docx] {title}: 去目录 {nt} 行, cjk={cjk_count(t):,}")
            text_to_docx(t, p, title)
        report[out] = {"kind": "docx", "cjk": cjk_count(t)}

    n_docs = 0
    with open(os.path.join(RAW, "gov_docs.jsonl"), encoding="utf-8") as f:
        for i, line in enumerate(f, 1):
            d = json.loads(line)
            safe = re.sub(r"[\\/:*?\"<>|\s]", "_", d["title"])[:50]
            p = os.path.join(GOV, f"gov_{i:03d}_{safe}.docx")
            if not os.path.exists(p):
                text_to_docx(d["text"], p, d["title"])
            n_docs += 1
    gov_cjk = 0
    with open(os.path.join(RAW, "gov_docs.jsonl"), encoding="utf-8") as f:
        for line in f:
            gov_cjk += json.loads(line)["cjk"]
    print(f"[docx] gov notices: {n_docs} files, cjk={gov_cjk:,}")
    report["gov/*.docx"] = {"kind": "docx", "files": n_docs, "cjk": gov_cjk}

    with open(os.path.join(PREP, "_prepare_report.json"), "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    total = sum(v.get("cjk", 0) for v in report.values()) + 30666
    print(f"\nprepared total (incl. stats 30,666): {total:,} chars")


if __name__ == "__main__":
    main()
